#!/usr/bin/python3

import docx

document = docx.Document()

document.add_heading('rocFFT Design Documentation', 0)


document.add_paragraph('''\
This section outlines the design of the rocFFT software library. A \
high-level description of the algorithms used in rocFFT is included \
along with various implementation-level details.''')

document.add_heading('Introduction', 1)

document.add_paragraph('''\
rocFFT is an open-source software library which is part of the \
Radeon Open Compute software ecosystem.  rocFFT is written in the HIP \
programming language and allows one to compute Fast Fourier Transforms \
(FFTs) on GPU devices.''')


document.add_heading('Algorithms', 1)

document.add_paragraph('This section describes the algorithms used in rocFFT.')

document.add_heading('Problem Lengths', 2)

document.add_paragraph('''\
The FFT algorithm computes the discrete Fourier transform by \
decomposing the problem into subproblems.  The most efficient FFT \
algorithms divide a problem of length N into small prime powers, which \
are referred to as radices.

The Stockham algorithm, a variant of the classic Tukey-Cooley \
algorithm, is used to compute FFTs for transforms on length N \
where N is decomposable into small prime factors.  rocFFT currently \
supports radix 2, 3, and 5: any transform whose length is composed of \
factors of 2, 3, and 5 can be computed in rocFFT using Stockham.

For other transform lengths, such as when N is a large prime, or when \
the decomposition of N includes a large prime, we make use of the \
Bluestein algorithm, which computes general-length transforms via \
convolution.  The convolution itself is computed via the convolution \
theorem, which states that a Fourier transform maps a cyclic \
convolution to a direct (Hadamard) product.  A linear convolution is \
achieved from the cyclic convolution extending the input by \
zero-padding.  This padding requires that the convolution length be at \
least 2N; moreover, we are free to choose a length greater than 2N, \
to, for example, the next binary power, at which point the Stockham \
method may be used.''')


document.add_heading('Real Data', 2)

document.add_paragraph('''\
The basic FFT method takes complex-valued input and produces \
complex-valued output, but many applications of the FFT use \
real-valued data.  For forward transforms with real input, the output \
is complex with Hermitian symmetry, and, making use of this symmetry, \
one need store approximately half of the output values.

The simplest method for FFTs on real data is to embed the transform \
into a complex transform of the same length where the imaginary part \
is set to zero.  One then performs a complex-to-complex transform and \
extracts the relevant data.

For even problem sizes, one can avoid having to use a complex buffer \
using one stage of the FFT to pre- or post-process the data to produce \
an N/2 complex transform.  This uses about half the memory and \
requires half as much computation as the complex-embedding method.''')


document.add_heading('Multi-dimensional Transforms', 2)

document.add_paragraph('''\
rocFFT handles 1D, 2D, and 3D transforms.  Multi-dimensional \
transforms are handled by performing batched one-dimensional transforms in the \
direction for which the data is contiguous.  After transforming in one \
direction, the data is transposed so that transforms can be performed \
in the next direction.''')

document.add_heading('Implementation', 1)

document.add_paragraph('''\
rocFFT's implementation of the above algorithms is split between \
host (CPU) and device (GPU) computation.  Host code is compiled into \
the shared libary librocfft.so, and device kernels are compiled into \
the shared library librocfft-devicec.so.''')

document.add_heading('Host Implementation', 2)

document.add_paragraph('''\
The information for performing a transform is stored in a plan, \
which is created by the host.  Plans contain a tree structure which \
represents the individual execution stages to be performed on the \
device.  Using a tree structure allows one to break the problem down \
into sub-problems, possibly using the same algorithm at different \
stages.  Input/output and work buffers, data types, and data layout \
are then assigned by traversing the tree structure.

The plan is executed by traversing the leaf nodes of the tree.  Each \
leaf contains a device kernel which is executed in order during plan \
execution.

In general, there are multiple methods for computing a specific \
transform, and the planner attempts to make the best choice based on \
an analysis of performance and memory requirements.

Some algorithms require extra work memory; these buffers are provided \
by the user at the execution stage.''')

document.add_paragraph('''\
The flow diagram below shows the implementation of rocFFT APIs and \
the connection to internal classes at a high level. All rocFFT APIs \
communicate to the rocFFT::Repo data structure. rocFFT::Repo is a \
singleton that manages a map of user plans to internal plans.  Each \
instance of rocFFT::execPlan holds a vector of rocFFT::TreeNode \
objects.''')
document.add_picture('flow.png', width=docx.shared.Inches(6))

document.add_heading('Device kernels and code generation', 2)

document.add_paragraph('''\
The majority of kernels are generated, with a few special cases \
written by hand.  Kernels are compiled at the same time as the rest of \
the C++ code; no compilation is done at run-time.  There are two basic \
types of kernels: computation kernels and data movement (transpose) \
kernels.

The kernel source file generator takes configurations from Generator \
Parameters and generates one-dimensional FFT kernels statically for \
all supported radices.''')

document.save('rocfft_documentation.docx')

